# ============================================================================
# app/report.py — the run report as a Word document
# ----------------------------------------------------------------------------
# Every task's run report is tables: settings, values the job ran with, the
# numbers it produced, and what each output file is. A .docx lays those out
# as real tables; a text file had to line them up by hand. This is the one
# place the document is styled, so the reports of different tasks read the
# same. Front-end agnostic, like the rest of app/: no gradio, no science.
# ============================================================================
from typing import Optional, Sequence

FONT_PT = 10
TABLE_FONT_PT = 9
MARGIN_LR_CM = 2.0
MARGIN_TB_CM = 1.8


def new_document():
    """An empty document in the app's report style."""
    from docx import Document
    from docx.shared import Cm, Pt
    document = Document()
    document.styles["Normal"].font.size = Pt(FONT_PT)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(MARGIN_LR_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_TB_CM)
    return document


def add_table(document, header: Sequence[str], rows: Sequence[Sequence],
              widths_cm: Optional[Sequence[float]] = None):
    """A grid table with a bold header row, in the smaller table font,
    followed by an empty paragraph so consecutive tables do not fuse."""
    from docx.shared import Cm, Pt
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
        run.font.size = Pt(TABLE_FONT_PT)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(text))
            run.font.size = Pt(TABLE_FONT_PT)
    if widths_cm:
        for column, width in zip(table.columns, widths_cm):
            for cell in column.cells:
                cell.width = Cm(width)
    document.add_paragraph()
    return table
